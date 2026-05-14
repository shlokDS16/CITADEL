"""
Text extraction from PDF / DOCX / images.

Strategy per file type:
  - DOCX  → python-docx (paragraphs + tables + headers).  Confidence = 99%.
  - PDF   → PyMuPDF (`fitz`).  If extracted text < 20 chars (likely scanned),
            caller falls back to OCR.
  - Image → caller goes straight to OCR.
"""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from typing import Optional

import fitz  # PyMuPDF
from docx import Document as DocxDocument

log = logging.getLogger("citadel.extractor")

DIGITAL_PDF_MIN_CHARS = 20  # fewer chars than this → assume scanned → OCR fallback


@dataclass
class ExtractionResult:
    text: str
    page_count: int
    confidence: float          # 0..100
    source: str                # "docx" | "pdf-digital" | "pdf-needs-ocr" | "image-needs-ocr" | "error"
    pages: list[str] = field(default_factory=list)   # per-page text (PDF only)
    error: Optional[str] = None


def extract_docx(content: bytes) -> ExtractionResult:
    """Read a .docx file and concatenate paragraphs + tables + headers."""
    try:
        doc = DocxDocument(io.BytesIO(content))
        chunks: list[str] = []

        for para in doc.paragraphs:
            if para.text.strip():
                chunks.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    chunks.append(row_text)

        for section in doc.sections:
            for header in (section.header, section.first_page_header):
                for para in header.paragraphs:
                    if para.text.strip():
                        chunks.append(para.text)

        text = "\n".join(chunks)
        return ExtractionResult(
            text=text,
            page_count=1,  # python-docx does not expose page count reliably
            confidence=99.0,
            source="docx",
        )
    except Exception as e:
        log.exception("DOCX extraction failed")
        return ExtractionResult(text="", page_count=0, confidence=0.0, source="error", error=str(e))


def extract_pdf_digital(content: bytes) -> ExtractionResult:
    """Try PyMuPDF text extraction. Returns hint about needing OCR fallback."""
    try:
        doc = fitz.open(stream=content, filetype="pdf")
        pages: list[str] = []
        for page in doc:
            pages.append(page.get_text("text"))
        doc.close()

        joined = "\n".join(pages).strip()
        if len(joined) < DIGITAL_PDF_MIN_CHARS:
            return ExtractionResult(
                text=joined,
                page_count=len(pages),
                confidence=0.0,
                source="pdf-needs-ocr",
                pages=pages,
            )

        return ExtractionResult(
            text=joined,
            page_count=len(pages),
            confidence=98.0,
            source="pdf-digital",
            pages=pages,
        )
    except Exception as e:
        log.exception("PDF extraction failed")
        return ExtractionResult(text="", page_count=0, confidence=0.0, source="error", error=str(e))


def split_pdf_pages(content: bytes, max_bytes_per_page: int) -> list[bytes]:
    """
    Split a PDF into per-page PDF byte blobs (so each page is < max_bytes_per_page).
    Used when sending to OCR.space (free tier 1MB/request).
    """
    src = fitz.open(stream=content, filetype="pdf")
    out: list[bytes] = []
    try:
        for i in range(src.page_count):
            single = fitz.open()
            single.insert_pdf(src, from_page=i, to_page=i)
            blob = single.tobytes()
            single.close()
            if len(blob) > max_bytes_per_page:
                # Single page still too big — render as PNG (often smaller than PDF)
                page = src.load_page(i)
                pix = page.get_pixmap(dpi=150)
                blob = pix.tobytes("png")
            out.append(blob)
    finally:
        src.close()
    return out


def render_last_page_png(content: bytes, dpi: int = 150) -> bytes:
    """Render the last page of a PDF to PNG bytes (used by signature.py)."""
    src = fitz.open(stream=content, filetype="pdf")
    try:
        page = src.load_page(src.page_count - 1)
        pix = page.get_pixmap(dpi=dpi)
        return pix.tobytes("png")
    finally:
        src.close()
